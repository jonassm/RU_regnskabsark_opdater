"""Selvst?ndig importmotor; kr?ver kun Python 3.10+ og de medf?lgende biblioteker."""
from pathlib import Path
import sys
sys.dont_write_bytecode = True
sys.path.insert(0, str(Path(__file__).resolve().parent / 'biblioteker.zip'))
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
import json
import os
import re
import shutil
import tempfile
import xml.etree.ElementTree as ET
from xml.sax.saxutils import escape
from zipfile import ZipFile, BadZipFile
import openpyxl
CELL_PATTERN = re.compile(r"<c\b[^>]*?(?:/>|>.*?</c>)", re.DOTALL)


class ImportErrorDetail(Exception):
    """A safe, user-facing validation error."""


def account_number(value):
    """API account numbers are integers; normalize numeric Excel IDs, including 0010."""
    if isinstance(value, bool) or value is None:
        raise ValueError("missing or boolean account number")
    if isinstance(value, (int, float, Decimal)):
        number = Decimal(str(value))
        if not number.is_finite() or number != number.to_integral_value():
            raise ValueError("non-integral account number")
        value = str(int(number))
    if not isinstance(value, str) or not re.fullmatch(r"[0-9]{1,9}", value.strip()):
        raise ValueError("account number must contain 1–9 digits")
    number = int(value.strip())
    if number < 1:
        raise ValueError("account number must be positive")
    return number


def select_column(sheet, year):
    if sheet.title != "Ny saldostruktur":
        raise ImportErrorDetail("The first sheet must be Ny saldostruktur, as in the inspected workbook.")
    if sheet["B1"].value != "Kontonr." or sheet["C1"].value != "Kontonavn":
        raise ImportErrorDetail("The expected Kontonr./Kontonavn headers have changed.")
    candidates = []
    for cell in sheet[1]:
        header = str(cell.value or "").lower()
        header = header.removeprefix("mock – ").removeprefix("delvis import – ")
        if re.search(rf"(?<!\d){year}(?!\d)", header) and header.startswith(
            ("foreløbigt", "resultat", "årsresultat", "bogført", "regnskab")
        ):
            candidates.append(cell.column)
    if len(candidates) != 1:
        raise ImportErrorDetail(
            f"Expected one actuals column for {year} in {sheet.title}; found {len(candidates)}. "
            "Prepare an actuals column manually; budgets are never used."
        )
    return candidates[0]


@dataclass(frozen=True)
class Target:
    account: int
    cell: str
    name: str
    multiplier: int


def build_account_mapping(sheet, column):
    grouped = defaultdict(list)
    issues = {"duplicate_excel_accounts": {}, "malformed_excel_accounts": [],
              "protected_formula_cells": [], "ambiguous_excel_rows": []}
    section = None
    for row in sheet.iter_rows(min_row=2):
        label = str(row[0].value or "").strip().upper()
        if label == "INDTÆGTER":
            section = "income"
        elif label == "UDGIFTER" or label.startswith("UDGIFTER "):
            section = "expense"
        account_cell = row[1]
        if account_cell.value is None:
            continue
        try:
            number = account_number(account_cell.value)
        except ValueError:
            issues["malformed_excel_accounts"].append(account_cell.coordinate)
            continue
        cell = sheet.cell(account_cell.row, column)
        grouped[number].append((cell, row[2].value, section))
    mapping = {}
    for number, rows in grouped.items():
        if len(rows) != 1:
            issues["duplicate_excel_accounts"][number] = [r[0].coordinate for r in rows]
            continue
        cell, name, section = rows[0]
        if cell.data_type == "f":
            issues["protected_formula_cells"].append(cell.coordinate)
        elif (section is None or not isinstance(name, str) or not name.strip()
              or any(cell.coordinate in merged for merged in sheet.merged_cells.ranges)
              or (cell.value is not None and not isinstance(cell.value, (int, float, Decimal)))
              or isinstance(cell.value, bool)):
            issues["ambiguous_excel_rows"].append(cell.coordinate)
        else:
            mapping[number] = Target(number, cell.coordinate, name, -1 if section == "income" else 1)
    return mapping, issues


def patch_sheet(xml, updates, formula_updates=None):
    """Retain original XML verbatim outside changed cells and cached formula results."""
    found = set()
    formula_updates = formula_updates or {}
    found_formulas = set()

    def replace(match):
        original = match.group()
        coord_match = re.search(r'\br="([A-Z]+[0-9]+)"', original)
        if not coord_match:
            return original
        coord = coord_match[1]
        if coord in formula_updates:
            change = formula_updates[coord]
            formula = re.search(r"<f>(.*?)</f>", original, flags=re.DOTALL)
            if not formula or "=" + formula[1] != change["old"] or coord in updates:
                raise ImportErrorDetail(f"Resultatformlen i {coord} matcher ikke den kontrollerede formel.")
            original = original[:formula.start()] + "<f>" + escape(change["new"][1:]) + "</f>" + original[formula.end():]
            found_formulas.add(coord)
            return re.sub(r"<v(?:\s[^>]*)?>.*?</v>|<v\s*/>", "", original, flags=re.DOTALL)
        if coord not in updates:
            # Excel recalculates these on opening; do not expose stale cached totals.
            if re.search(r"<f(?:\s|>)", original):
                return re.sub(r"<v(?:\s[^>]*)?>.*?</v>|<v\s*/>", "", original, flags=re.DOTALL)
            return original
        if "<f" in original:
            raise ImportErrorDetail(f"Refusing to overwrite formula {coord}.")
        found.add(coord)
        opening = original[:original.index(">")].rstrip("/")
        opening = re.sub(r'\s+t="[^"]*"', "", opening)
        body = "" if original.endswith("/>") else original[original.index(">") + 1:-4]
        body = re.sub(r"<(v|is)\b[^>]*>.*?</\1>|<(?:v|is)\s*/>", "", body, flags=re.DOTALL)
        value = updates[coord]
        if isinstance(value, str):
            return opening + ' t="inlineStr"><is><t>' + escape(value) + "</t></is>" + body + "</c>"
        if not isinstance(value, (int, float, Decimal)) or not Decimal(str(value)).is_finite():
            raise ImportErrorDetail(f"Invalid numeric value at {coord}.")
        return opening + "><v>" + str(value) + "</v>" + body + "</c>"

    result = CELL_PATTERN.sub(replace, xml)
    if found != set(updates):
        raise ImportErrorDetail("Target cells are missing from the original XML; no file saved.")
    if found_formulas != set(formula_updates):
        raise ImportErrorDetail("Den forventede resultatformel mangler; ingen fil gemt.")
    ET.fromstring(result)
    return result

def parse_amount(value):
    if isinstance(value, bool) or value is None:
        raise ValueError("Beløbet mangler eller er ugyldigt")
    if isinstance(value, str):
        text = value.strip().replace("\u2212", "-").replace("\u00a0", "").replace("\u202f", "")
        # Text cells follow Danish notation; ordinary numeric Excel cells need no conversion.
        if not re.fullmatch(r"[+-]?(?:[0-9]+|[0-9]{1,3}(?:\.[0-9]{3})+)(?:,[0-9]{1,2})?", text):
            raise ValueError("Beløbet er ikke et dansk tal")
        text = text.replace(".", "").replace(",", ".")
    elif isinstance(value, (int, float, Decimal)):
        text = str(value)
    else:
        raise ValueError("Ugyldig beløbstype")
    try:
        amount = Decimal(text)
    except InvalidOperation:
        raise ValueError("Ugyldigt beløb") from None
    if not amount.is_finite() or abs(amount) >= Decimal("1e13"):
        raise ValueError("Beløbet overskrider sikker Excel-præcision")
    return amount

def render_workbook(source_bytes, sheet_name, updates, formula_updates=None):
    """Copy all ZIP parts unchanged except target worksheet and recalculation flags."""
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    with ZipFile(BytesIO(source_bytes)) as source:
        workbook_xml = source.read("xl/workbook.xml").decode("utf-8")
        root = ET.fromstring(workbook_xml)
        sheet = next(s for s in root.find("m:sheets", ns) if s.attrib["name"] == sheet_name)
        rel_id = sheet.attrib[f"{{{rel_ns}}}id"]
        rels = ET.fromstring(source.read("xl/_rels/workbook.xml.rels"))
        target = next(r.attrib["Target"] for r in rels if r.attrib["Id"] == rel_id)
        part = target.lstrip("/") if target.startswith("/") else "xl/" + target
        changed_sheet = patch_sheet(source.read(part).decode("utf-8"), updates, formula_updates).encode("utf-8")
        calc = '<calcPr calcMode="auto" fullCalcOnLoad="1" forceFullCalc="1"/>'
        if re.search(r"<calcPr\b", workbook_xml):
            workbook_xml = re.sub(r"<calcPr\b[^>]*?(?:/>|>.*?</calcPr>)", calc, workbook_xml)
        else:
            raise ImportErrorDetail("Expected workbook calculation settings are missing.")
        output_bytes = BytesIO()
        with ZipFile(output_bytes, "w") as output:
            output.comment = source.comment
            for info in source.infolist():
                data = changed_sheet if info.filename == part else (
                    workbook_xml.encode("utf-8") if info.filename == "xl/workbook.xml" else source.read(info))
                output.writestr(info, data)
    # Validate the completed in-memory file before publishing any output.
    check = openpyxl.load_workbook(BytesIO(output_bytes.getvalue()), data_only=False)
    expected_values = dict(updates)
    expected_values.update({coord: change["new"] for coord, change in (formula_updates or {}).items()})
    for coordinate, value in expected_values.items():
        actual = check[sheet_name][coordinate].value
        if isinstance(value, str):
            valid = actual == value
        else:
            valid = actual is not None and abs(Decimal(str(actual)) - Decimal(str(value))) <= Decimal("0.0000001")
        if not valid:
            raise ImportErrorDetail(f"Saved value failed verification at {coordinate}.")
    check.close()
    return output_bytes.getvalue()
